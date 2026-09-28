#!/usr/bin/env python3
"""PDF -> PPTX, one slide per page, each page placed full-bleed as an image.

This is the route that produces a deck a colleague can open in PowerPoint while
keeping the LaTeX typesetting exactly as rendered. The trade is that the text is
not selectable or editable - each slide is a picture. Slidev's own PPTX export has
the same property, so this is not a regression against the current toolchain.

Usage: python3 pdf2pptx.py <in.pdf> <out.pptx> [dpi]
"""
import subprocess, sys, tempfile
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

EMU_PER_IN = 914400


def main(pdf: Path, out: Path, dpi: int = 200) -> None:
    with tempfile.TemporaryDirectory() as td:
        subprocess.run(["pdftoppm", "-png", "-r", str(dpi), str(pdf), f"{td}/pg"],
                       check=True)
        pages = sorted(Path(td).glob("pg-*.png"))
        if not pages:
            raise SystemExit(f"no pages rendered from {pdf}")

        from PIL import Image
        w_px, h_px = Image.open(pages[0]).size
        prs = Presentation()
        prs.slide_width = Emu(round(w_px / dpi * EMU_PER_IN))
        prs.slide_height = Emu(round(h_px / dpi * EMU_PER_IN))
        blank = prs.slide_layouts[6]            # 6 is the blank layout
        for p in pages:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(p), 0, 0,
                                     width=prs.slide_width, height=prs.slide_height)
        prs.save(str(out))
        print(f"{out}  —  {len(pages)} slide(s) at {w_px}x{h_px} px, {dpi} dpi")


if __name__ == "__main__":
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]), Path(sys.argv[2]),
         int(sys.argv[3]) if len(sys.argv) > 3 else 200)
